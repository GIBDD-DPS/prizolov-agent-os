# Prizolov Agent OS 0.3.0 | Author: Dm.Andreyanov | Brand: Prizolov Lab | © 2026
# SPDX-FileCopyrightText: 2026 Dm.Andreyanov / Prizolov Lab
# SPDX-License-Identifier: Apache-2.0

"""Тесты чтения PDF, Word, Excel и построения графиков."""

import json
import math
import random
from datetime import date, datetime, timedelta

import docx
import matplotlib.pyplot as plt
import openpyxl
import pytest
from PIL import Image
from pypdf import PdfWriter

from prizolov_os.__about__ import PROJECT_ID, SIGNATURE
from prizolov_os.analytics import parse_cashflow_csv
from prizolov_os.charts import daily_balance
from prizolov_os.documents import DocumentError, extract_text, sheet_as_csv
from prizolov_os.forecasting import ForecastEngine
from prizolov_os.llm import ToolCall
from prizolov_os.market import PriceSeries
from prizolov_os.security import unwrap
from prizolov_os.tools import ToolRegistry, Workspace, cashflow_tool, chart_tools, default_tools


def make_xlsx(path, rows, title="Выписка"):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    workbook.create_sheet("Второй").append(["заметка"])
    workbook.save(path)


def make_pdf(path, text):
    fig = plt.figure()
    fig.text(0.1, 0.5, text)
    fig.savefig(path, format="pdf")
    plt.close(fig)


BANK_ROWS = [
    ["Дата", "Сумма", "Назначение"],
    [datetime(2026, 9, 1), 150000.0, "Оплата от клиента"],
    [datetime(2026, 9, 5), -40000, "Аренда"],
    [datetime(2026, 9, 20), -30000.5, "Закупки"],
]


class TestDocuments:
    def test_docx_paragraphs_and_tables(self, tmp_path):
        document = docx.Document()
        document.add_paragraph("Договор поставки № 15")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Товар", "Цена"
        table.cell(1, 0).text, table.cell(1, 1).text = "Бумага", "500"
        document.save(tmp_path / "d.docx")
        text = extract_text(tmp_path / "d.docx")
        assert "Договор поставки № 15" in text
        assert "Бумага | 500" in text

    def test_xlsx_all_sheets(self, tmp_path):
        make_xlsx(tmp_path / "b.xlsx", BANK_ROWS)
        text = extract_text(tmp_path / "b.xlsx")
        assert "## Лист: Выписка" in text and "## Лист: Второй" in text
        assert "2026-09-01\t150000\tОплата от клиента" in text

    def test_sheet_as_csv_feeds_cashflow_parser(self, tmp_path):
        make_xlsx(tmp_path / "b.xlsx", BANK_ROWS)
        transactions = parse_cashflow_csv(sheet_as_csv(tmp_path / "b.xlsx"))
        assert [t.amount for t in transactions] == [150000, -40000, -30000.5]
        assert transactions[0].date == date(2026, 9, 1)

    def test_missing_sheet(self, tmp_path):
        make_xlsx(tmp_path / "b.xlsx", BANK_ROWS)
        with pytest.raises(DocumentError, match="Нет листа"):
            sheet_as_csv(tmp_path / "b.xlsx", sheet="Нет такого")

    def test_pdf_text(self, tmp_path):
        make_pdf(tmp_path / "r.pdf", "Quarterly revenue 1500")
        text = extract_text(tmp_path / "r.pdf")
        assert "--- страница 1 ---" in text and "Quarterly revenue 1500" in text

    def test_scanned_pdf_reported(self, tmp_path):
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        with open(tmp_path / "scan.pdf", "wb") as f:
            writer.write(f)
        with pytest.raises(DocumentError, match="OCR"):
            extract_text(tmp_path / "scan.pdf")

    def test_corrupted_file(self, tmp_path):
        (tmp_path / "bad.docx").write_bytes(b"not a zip")
        with pytest.raises(DocumentError, match="Не удалось прочитать"):
            extract_text(tmp_path / "bad.docx")

    def test_read_file_tool_reads_documents_as_untrusted(self, tmp_path):
        make_xlsx(tmp_path / "b.xlsx", BANK_ROWS)
        result = ToolRegistry(default_tools(tmp_path)).execute(
            ToolCall("t", "read_file", {"path": "b.xlsx"})
        )
        assert not result.is_error
        assert "Оплата от клиента" in unwrap(result.output)
        assert result.output.startswith("<untrusted_data")

    def test_cashflow_tool_accepts_excel(self, tmp_path):
        make_xlsx(tmp_path / "b.xlsx", BANK_ROWS)
        result = ToolRegistry([cashflow_tool(Workspace(tmp_path))]).execute(
            ToolCall("t", "analyze_cashflow",
                     {"path": "b.xlsx", "opening_balance": 0, "horizon_days": 30})
        )
        assert json.loads(unwrap(result.output))["closing_balance"] == pytest.approx(79999.5)


def prices(n=200, seed=3):
    rng = random.Random(seed)
    values = [100.0]
    for _ in range(n):
        values.append(values[-1] * math.exp(rng.gauss(0.0005, 0.01)))
    return [date(2026, 1, 1) + timedelta(i) for i in range(len(values))], values


class StubMarket:
    def history(self, source, symbol, days):
        dates, values = prices()
        return PriceSeries(source, symbol, "USD", dates, values)


def png_info(path):
    with Image.open(path) as image:
        return image.size, dict(image.info)


class TestCharts:
    def registry(self, tmp_path):
        return ToolRegistry(chart_tools(StubMarket(), Workspace(tmp_path), ForecastEngine()))

    def test_market_chart(self, tmp_path):
        result = self.registry(tmp_path).execute(ToolCall("t", "chart_market", {
            "source": "yahoo", "symbol": "GC=F", "history_days": 200, "horizon_days": 30,
        }))
        assert not result.is_error, result.output
        path = tmp_path / json.loads(result.output)["chart"]
        assert path.parent.name == "charts" and path.suffix == ".png"
        size, info = png_info(path)
        assert size[0] > 800
        assert info["Author"] == "Dm.Andreyanov / Prizolov Lab"
        assert PROJECT_ID in info["Description"] and SIGNATURE in info["Description"]

    def test_market_chart_from_file(self, tmp_path):
        dates, values = prices()
        lines = ["date;close"] + [f"{d.isoformat()};{v}" for d, v in zip(dates, values)]
        (tmp_path / "p.csv").write_text("\n".join(lines))
        result = self.registry(tmp_path).execute(ToolCall("t", "chart_market", {
            "source": "file", "symbol": "p.csv", "history_days": 0, "horizon_days": 14,
        }))
        assert not result.is_error, result.output

    def test_cashflow_charts(self, tmp_path):
        make_xlsx(tmp_path / "b.xlsx", BANK_ROWS)
        result = self.registry(tmp_path).execute(ToolCall("t", "chart_cashflow", {
            "path": "b.xlsx", "opening_balance": 10000, "horizon_days": 30,
        }))
        assert not result.is_error, result.output
        charts = json.loads(result.output)["charts"]
        assert len(charts) == 2
        for chart in charts:
            assert png_info(tmp_path / chart)[1]["Copyright"].startswith("©")

    def test_chart_does_not_record_forecast(self, tmp_path):
        from prizolov_os.forecasting import ForecastJournal
        from prizolov_os.memory import Store

        journal = ForecastJournal(Store())
        registry = ToolRegistry(
            chart_tools(StubMarket(), Workspace(tmp_path), ForecastEngine(journal))
        )
        registry.execute(ToolCall("t", "chart_market", {
            "source": "cbr", "symbol": "USD", "history_days": 200, "horizon_days": 30,
        }))
        assert journal.counts() == {}

    def test_daily_balance(self):
        transactions = parse_cashflow_csv("дата;сумма\n01.09.2026;100\n03.09.2026;-30\n")
        days, balances = daily_balance(transactions, 50)
        assert days == [date(2026, 9, 1), date(2026, 9, 2), date(2026, 9, 3)]
        assert balances == [150, 150, 120]
