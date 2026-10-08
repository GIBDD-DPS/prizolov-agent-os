# Prizolov Agent OS 0.3.0 | Author: Dm.Andreyanov | Brand: Prizolov Lab | © 2026
# SPDX-FileCopyrightText: 2026 Dm.Andreyanov / Prizolov Lab
# SPDX-License-Identifier: Apache-2.0

"""Тесты базы знаний (RAG)."""

import io
import os

import docx
import matplotlib.pyplot as plt
import openpyxl
import pytest
from rich.console import Console

from cli.app import ChatApp
from prizolov_os.core.kernel import Kernel
from prizolov_os.knowledge import KnowledgeBase, build_match, split_text
from prizolov_os.llm import FakeLLMClient, ToolCall, tool_use_response
from prizolov_os.memory import Store
from prizolov_os.security import unwrap
from prizolov_os.tools import ToolRegistry, knowledge_tool


@pytest.fixture
def docs(tmp_path):
    (tmp_path / "policy.md").write_text(
        "# Политика скидок\n\nКлиентам с оборотом от 1 млн ₽ даём скидку 5%.\n\n"
        "Оплата по договору в течение 10 банковских дней."
    )
    (tmp_path / "gold.txt").write_text("Золото подорожало за квартал на 12 процентов.")
    document = docx.Document()
    document.add_paragraph("Договор аренды склада. Арендная плата 90 000 рублей в месяц.")
    document.save(tmp_path / "rent.docx")
    workbook = openpyxl.Workbook()
    workbook.active.title = "Тарифы"
    workbook.active.append(["Доставка Москва", 500])
    workbook.save(tmp_path / "prices.xlsx")
    fig = plt.figure()
    fig.text(0.1, 0.5, "Warranty period is 24 months")
    fig.savefig(tmp_path / "warranty.pdf", format="pdf")
    plt.close(fig)
    (tmp_path / "charts").mkdir()
    (tmp_path / "charts" / "note.txt").write_text("служебный файл графиков")
    (tmp_path / "image.png").write_bytes(b"\x89PNG")
    return tmp_path


def kb(root):
    return KnowledgeBase(Store(), root)


class TestIndex:
    def test_indexes_supported_files_only(self, docs):
        base = kb(docs)
        report = base.index()
        assert sorted(report.added) == [
            "gold.txt", "policy.md", "prices.xlsx", "rent.docx", "warranty.pdf"
        ]
        assert base.stats()["files"] == 5

    def test_incremental_update_and_removal(self, docs):
        base = kb(docs)
        base.index()
        assert not base.index().changed
        (docs / "gold.txt").write_text("Серебро подешевело.")
        os.utime(docs / "gold.txt", (1, 1))
        (docs / "policy.md").unlink()
        report = base.index()
        assert report.updated == ["gold.txt"] and report.removed == ["policy.md"]
        assert base.search("политика скидок") == []
        assert base.search("серебро")[0].path == "gold.txt"

    def test_broken_file_reported(self, docs):
        (docs / "broken.docx").write_bytes(b"not a zip")
        base = kb(docs)
        report = base.index()
        assert report.errors and report.errors[0][0] == "broken.docx"
        assert base.stats()["errors"] == 1


class TestSearch:
    @pytest.mark.parametrize("query, path", [
        ("какая скидка крупным клиентам", "policy.md"),
        ("цена золота", "gold.txt"),
        ("арендная плата за склад", "rent.docx"),
        ("стоимость доставки", "prices.xlsx"),
        ("warranty months", "warranty.pdf"),
    ])
    def test_finds_by_word_forms(self, docs, query, path):
        assert kb(docs).search(query)[0].path == path

    def test_locations(self, docs):
        base = kb(docs)
        assert base.search("доставка")[0].source == "prices.xlsx, лист «Тарифы»"
        assert base.search("warranty")[0].location == "стр. 1"

    def test_no_match_and_empty_query(self, docs):
        base = kb(docs)
        assert base.search("квантовая хромодинамика") == []
        assert base.search("?!") == []

    def test_match_syntax_is_safe(self):
        assert build_match('скидка "OR" NOT*') == '"скидк"* OR "or"* OR "not"*'


def test_split_text_respects_size():
    chunks = split_text(("слово " * 400 + "\n\n") * 3)
    assert len(chunks) > 3
    assert all(len(c) <= 1300 for c in chunks)
    assert split_text("короткий текст") == ["короткий текст"]


class TestTool:
    def test_output_cites_sources_and_is_untrusted(self, docs):
        result = ToolRegistry([knowledge_tool(kb(docs))]).execute(
            ToolCall("t", "search_knowledge", {"query": "скидка"})
        )
        text = unwrap(result.output)
        assert result.output.startswith("<untrusted_data")
        assert "[1] policy.md" in text and "скидку 5%" in text

    def test_nothing_found(self, docs):
        result = ToolRegistry([knowledge_tool(kb(docs))]).execute(
            ToolCall("t", "search_knowledge", {"query": "космос"})
        )
        assert "ничего не найдено" in unwrap(result.output)


class TestKernelAndCli:
    def test_agents_get_search(self, docs):
        kernel = Kernel.create(llm=FakeLLMClient(), store=Store(), workspace_dir=docs)
        with_search = {n for n, a in kernel.agents.items() if "search_knowledge" in a.tools}
        assert with_search == {"director", "researcher", "writer"}

    def test_director_answers_from_documents(self, docs):
        kernel = Kernel.create(llm=FakeLLMClient([
            tool_use_response("search_knowledge", {"query": "скидка"}),
            "Скидка 5% (policy.md)",
        ]), store=Store(), workspace_dir=docs, self_check="off")
        result = kernel.chat("Какая у нас скидка?")
        assert result.text == "Скидка 5% (policy.md)"
        assert "policy.md" in result.tool_results[0].output

    def test_index_command(self, docs):
        kernel = Kernel.create(llm=FakeLLMClient(), store=Store(), workspace_dir=docs)
        out = io.StringIO()
        app = ChatApp(kernel, Console(file=out, width=150, color_system=None), lambda _: "")
        app.handle("/index")
        app.handle("/index")
        text = out.getvalue()
        assert "База знаний обновлена: новых 5 файлов" in text
        assert "Изменений в документах нет" in text
        assert "файлов 5" in text
