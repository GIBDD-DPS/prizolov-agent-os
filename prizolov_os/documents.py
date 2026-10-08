# Prizolov Agent OS 0.3.0 | Author: Dm.Andreyanov | Brand: Prizolov Lab | © 2026
# SPDX-FileCopyrightText: 2026 Dm.Andreyanov / Prizolov Lab
# SPDX-License-Identifier: Apache-2.0

"""Чтение документов: PDF, Word и Excel превращаются в текст или таблицу."""

import csv
import io
from datetime import date, datetime
from pathlib import Path
from typing import Any, List, Optional

DOCUMENT_SUFFIXES = {".pdf", ".docx", ".xlsx", ".xlsm"}
TABLE_SUFFIXES = {".xlsx", ".xlsm"}
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_TEXT_CHARS = 200_000
MAX_SHEET_ROWS = 5_000


class DocumentError(ValueError):
    """Документ не удалось прочитать."""


def is_document(path: Path) -> bool:
    return path.suffix.lower() in DOCUMENT_SUFFIXES


def extract_text(path: Path) -> str:
    """Текст документа. Длинные документы обрезаются с пометкой."""
    _check_size(path)
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            text = _pdf_text(path)
        elif suffix == ".docx":
            text = _docx_text(path)
        elif suffix in TABLE_SUFFIXES:
            text = _xlsx_text(path)
        else:
            raise DocumentError(f"Формат {suffix} не поддерживается")
    except DocumentError:
        raise
    except Exception as e:  # noqa: BLE001 - повреждённый или защищённый файл
        raise DocumentError(f"Не удалось прочитать {path.name}: {e}") from None
    if len(text) > MAX_TEXT_CHARS:
        note = f"\n\n[Показаны первые {MAX_TEXT_CHARS} символов из {len(text)}]"
        text = text[:MAX_TEXT_CHARS] + note
    return text


def sheet_as_csv(path: Path, sheet: Optional[str] = None) -> str:
    """Лист Excel как CSV с разделителем ';' (для анализа выписок и котировок)."""
    _check_size(path)
    workbook = _open_workbook(path)
    try:
        if sheet:
            if sheet not in workbook.sheetnames:
                raise DocumentError(f"Нет листа '{sheet}'. Листы: {', '.join(workbook.sheetnames)}")
            worksheet = workbook[sheet]
        else:
            worksheet = workbook.worksheets[0]
        buffer = io.StringIO()
        writer = csv.writer(buffer, delimiter=";")
        for row in _rows(worksheet):
            writer.writerow(row)
        return buffer.getvalue()
    finally:
        workbook.close()


def _check_size(path: Path) -> None:
    size = path.stat().st_size
    if size > MAX_FILE_BYTES:
        raise DocumentError(f"Файл слишком большой ({size // 1024 // 1024} МБ, максимум 50)")


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:  # noqa: BLE001
            raise DocumentError("PDF защищён паролем") from None
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"--- страница {number} ---\n{text}")
    if not pages:
        raise DocumentError(
            f"В {path.name} нет текстового слоя - вероятно, это скан. "
            "Распознавание сканов (OCR) не поддерживается."
        )
    return "\n\n".join(pages)


def _docx_text(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts: List[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for number, table in enumerate(document.tables, start=1):
        parts.append(f"\n[Таблица {number}]")
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _xlsx_text(path: Path) -> str:
    workbook = _open_workbook(path)
    try:
        parts = []
        for worksheet in workbook.worksheets:
            rows = ["\t".join(row) for row in _rows(worksheet)]
            parts.append(f"## Лист: {worksheet.title}\n" + "\n".join(rows))
        return "\n\n".join(parts)
    finally:
        workbook.close()


def _open_workbook(path: Path) -> Any:
    import openpyxl

    try:
        return openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise DocumentError(f"Не удалось открыть {path.name}: {e}") from None


def _rows(worksheet: Any) -> List[List[str]]:
    rows = []
    for row in worksheet.iter_rows(values_only=True):
        cells = [_cell(value) for value in row]
        while cells and not cells[-1]:
            cells.pop()
        if cells:
            rows.append(cells)
        if len(rows) >= MAX_SHEET_ROWS:
            break
    return rows


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == datetime.min.time():
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)
